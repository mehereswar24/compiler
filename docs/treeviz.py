"""Helpers that draw the parse tree and token chips as native PowerPoint shapes."""
from pptx.util import Inches as PI, Pt as PP
from pptx.dml.color import RGBColor as PC
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

MARK = ("|-- ", "`-- ")


def parse_statements(tree_lines):
    """tree_lines: lines after 'Program'. Returns a list of root nodes, one per statement."""
    roots, stack = [], []
    for line in tree_lines:
        idx = min((line.find(m) for m in MARK if m in line), default=-1)
        if idx < 0:
            continue
        node = {"label": line[idx + 4:], "kids": [], "depth": idx // 4}
        while stack and stack[-1]["depth"] >= node["depth"]:
            stack.pop()
        if stack:
            stack[-1]["kids"].append(node)
        else:
            roots.append(node)
        stack.append(node)
    return roots


def _layout(n, depth, ctr):
    n["d"] = depth
    if not n["kids"]:
        n["x"] = ctr[0]; ctr[0] += 1
    else:
        for k in n["kids"]:
            _layout(k, depth + 1, ctr)
        n["x"] = (n["kids"][0]["x"] + n["kids"][-1]["x"]) / 2


def _max_depth(n):
    return n["d"] if not n["kids"] else max(_max_depth(k) for k in n["kids"])


def _colour(label):
    w = label.split()[0]
    if w in ("Decl", "Assign", "Print"): return PC(0x6A, 0x3F, 0xA0)
    if w in ("BinOp", "Neg"): return PC(0x1B, 0x8A, 0x8A)
    if w in ("Int", "Float"): return PC(0x8A, 0x6A, 0x1F)
    if w == "Var": return PC(0x2C, 0x6E, 0x91)
    return PC(0x55, 0x5B, 0x66)


def draw_tree(slide, root, x, y, w, h, font=12):
    """Draw one statement tree inside the box (x, y, w, h) in inches."""
    ctr = [0]
    _layout(root, 0, ctr)
    leaves, levels = ctr[0], _max_depth(root) + 1
    cell_w = w / max(leaves, 1)
    cell_h = h / levels
    bw = min(cell_w * 0.92, 1.9)
    bh = min(cell_h * 0.62, 0.5)

    def centre(n):
        return x + (n["x"] + 0.5) * cell_w, y + n["d"] * cell_h + bh / 2

    def lines(n):
        cx, cy = centre(n)
        for k in n["kids"]:
            kx, ky = centre(k)
            c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, PI(cx), PI(cy + bh / 2), PI(kx), PI(ky - bh / 2))
            c.line.color.rgb = PC(0x9A, 0xA5, 0xB5); c.line.width = PP(1.5)
            lines(k)

    def boxes(n):
        cx, cy = centre(n)
        b = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PI(cx - bw / 2), PI(cy - bh / 2), PI(bw), PI(bh))
        b.fill.solid(); b.fill.fore_color.rgb = _colour(n["label"]); b.line.fill.background()
        tf = b.text_frame; tf.margin_left = tf.margin_right = PI(0.03); tf.margin_top = tf.margin_bottom = PI(0.01)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = False
        tf.text = n["label"]
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        p.font.size = PP(font); p.font.bold = True; p.font.color.rgb = PC(255, 255, 255)
        for k in n["kids"]:
            boxes(k)

    lines(root)
    boxes(root)


def _chip_colour(cls):
    if cls == "KEYWORD": return PC(0x6A, 0x3F, 0xA0)
    if cls == "IDENTIFIER": return PC(0x2C, 0x6E, 0x91)
    if cls.endswith("LITERAL"): return PC(0x8A, 0x6A, 0x1F)
    if cls in ("PLUS", "MINUS", "STAR", "SLASH", "PERCENT", "CARET"): return PC(0x1B, 0x8A, 0x8A)
    return PC(0x55, 0x5B, 0x66)


LEGEND = [("KEYWORD", "keyword"), ("IDENTIFIER", "identifier"), ("INT_LITERAL", "number"), ("PLUS", "operator"), ("ASSIGN", "punctuation / =")]


def draw_token_chips(slide, tokens, x, y, w, row_h=0.62, font=14):
    """tokens: list of (cls, lexeme, line). One row per source line."""
    rows = {}
    for cls, lex, line in tokens:
        rows.setdefault(line, []).append((cls, lex))
    for r, line in enumerate(sorted(rows)):
        ry = y + r * row_h
        t = slide.shapes.add_textbox(PI(x), PI(ry), PI(0.7), PI(0.45)); t.text_frame.text = f"{line}"
        t.text_frame.paragraphs[0].font.size = PP(13); t.text_frame.paragraphs[0].font.color.rgb = PC(0x88, 0x92, 0xA3)
        cx = x + 0.7
        for cls, lex in rows[line]:
            cw = max(0.45, 0.115 * len(lex) + 0.3)
            if cx + cw > x + w: break
            b = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PI(cx), PI(ry), PI(cw), PI(0.45))
            b.fill.solid(); b.fill.fore_color.rgb = _chip_colour(cls); b.line.fill.background()
            tf = b.text_frame; tf.margin_left = tf.margin_right = PI(0.02); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.text = lex; p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
            p.font.size = PP(font); p.font.bold = True; p.font.name = "Consolas"; p.font.color.rgb = PC(255, 255, 255)
            cx += cw + 0.1
    ly = y + len(rows) * row_h + 0.15
    lx = x + 0.7
    for cls, name in LEGEND:
        b = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PI(lx), PI(ly), PI(0.3), PI(0.3))
        b.fill.solid(); b.fill.fore_color.rgb = _chip_colour(cls); b.line.fill.background()
        t = slide.shapes.add_textbox(PI(lx + 0.35), PI(ly - 0.03), PI(1.9), PI(0.4)); t.text_frame.text = name
        t.text_frame.paragraphs[0].font.size = PP(13); t.text_frame.paragraphs[0].font.color.rgb = PC(0x55, 0x5B, 0x66)
        lx += 2.3
