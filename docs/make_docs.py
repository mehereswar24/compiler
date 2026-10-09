"""Generate REPORT.docx and PRESENTATION.pptx for the Review-1 (25%) scope.
All compiler output shown is captured live from arithc.exe."""
import subprocess, pathlib
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pptx import Presentation
from pptx.util import Inches as PI, Pt as PP
from pptx.dml.color import RGBColor as PC
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from treeviz import draw_tree, draw_token_chips, parse_statements

ROOT = pathlib.Path(__file__).resolve().parent.parent
OUT = ROOT / "docs"
EXE = str(ROOT / "arithc.exe")

def run(test, *flags):
    r = subprocess.run([EXE, *flags, str(ROOT / "tests" / test)], capture_output=True, text=True)
    return r.stdout.replace("\r", "").rstrip()

def section(text, start, end=None):
    i = text.index(start)
    j = text.index(end, i + 1) if end else len(text)
    return text[i:j].rstrip()

demo_src = (ROOT / "tests/t4_demo.ac").read_text().strip()
prec_src = (ROOT / "tests/t1_valid.ac").read_text().strip()
bad_src = (ROOT / "tests/t3_syntax.ac").read_text().strip()
demo_full = run("t4_demo.ac")
demo_tokens = section(demo_full, "== Token stream ==", "== Syntax").split("\n")[1:]
demo_tree = "\n".join(section(demo_full, "== Syntax (parse) tree ==", "== Summary").split("\n")[1:])
prec_q = run("t1_valid.ac", "-q")
prec_tree = "\n".join(section(prec_q, "== Syntax (parse) tree ==", "== Summary").split("\n")[1:])
p_lines = prec_tree.split("\n")
i_p = next(i for i, l in enumerate(p_lines) if "Decl int p" in l)
pow_tree = "\n".join(p_lines[i_p:i_p + 7])
bad_out = run("t3_syntax.ac", "-q")
bad_diag = section(bad_out, "== Phase 1+2", "== Syntax")
tests_out = subprocess.run(["sh", str(ROOT / "run_tests.sh")], capture_output=True, text=True, cwd=ROOT).stdout.replace("\r", "").rstrip()
n_pass = tests_out.count("PASS"); n_fail = tests_out.count("FAIL")

TITLE = "Compiler for an Arithmetic Expression Language using Lex/Flex and Yacc/Bison"
AUTHOR = "Metlapalli Mehereswar (23BAI0138)  |  Team members: <add names / reg. nos.>"
COURSE = "B.Tech CSE, VIT Vellore  |  Compiler Design  |  Faculty: <add name>"

GRAMMAR = """program -> stmts
stmts   -> stmts stmt | (empty)
stmt    -> type ID ';'
         | type ID '=' expr ';'
         | ID '=' expr ';'
         | 'print' expr ';'
         | error ';'                 (panic-mode recovery)
type    -> 'int' | 'float'
expr    -> expr ('+'|'-'|'*'|'/'|'%'|'^') expr
         | '-' expr | '+' expr | '(' expr ')'
         | INUM | FNUM | ID

Precedence (low -> high):   + -   <   * / %   <   unary -   <   ^
Associativity:              left for + - * / %,  right for ^"""

TOKENS = [
    ("KEYWORD", "int  float  print", "Reserved words"),
    ("IDENTIFIER", "[A-Za-z][A-Za-z0-9_]*", "Variable names"),
    ("INT_LITERAL", "[0-9]+", "Integer constants"),
    ("FLOAT_LITERAL", "1.5   .5   2.   1e3   2.5E-2", "Real constants, optional exponent"),
    ("PLUS MINUS STAR SLASH PERCENT CARET", "+  -  *  /  %  ^", "Arithmetic operators"),
    ("ASSIGN", "=", "Assignment"),
    ("LPAREN RPAREN SEMICOLON", "(  )  ;", "Punctuation"),
    ("(skipped)", "whitespace, // line, /* block */", "Discarded by the lexer"),
]

PLAN = [
    ("1", "Language definition, grammar, token set", "Done - Review 1"),
    ("2", "Flex lexer: token stream, line numbers, lexical errors", "Done - Review 1"),
    ("3", "Bison parser: precedence, parse-tree construction, error recovery", "Done - Review 1"),
    ("4", "Symbol table and semantic validation (declarations, type checking, int/float rules, division by zero)", "Planned - next review"),
    ("5", "Three Address Code generation (temporaries, type conversions, constant folding)", "Planned - next review"),
    ("6", "Testing of semantic and code-generation phases; final report and demo", "Planned - final review"),
]

LIT = [
    ("[1]", "Aho, Lam, Sethi, Ullman", "Compilers: Principles, Techniques, and Tools, 2nd ed., Pearson, 2006",
     "Standard reference for the phase structure, syntax-directed translation, three-address code and type checking."),
    ("[2]", "M. E. Lesk, E. Schmidt", "Lex - A Lexical Analyzer Generator, Bell Labs CSTR 39, 1975",
     "Original Lex: regular expressions compiled into a DFA-driven scanner. Basis of Flex."),
    ("[3]", "S. C. Johnson", "Yacc: Yet Another Compiler-Compiler, Bell Labs CSTR 32, 1975",
     "Original Yacc: LALR(1) parser generator with semantic actions. Basis of Bison."),
    ("[4]", "J. Levine", "flex & bison, O'Reilly, 2009",
     "Practical Flex/Bison integration (yylval, %union, token sharing, error recovery)."),
    ("[5]", "R. W. Floyd", "Syntactic analysis and operator precedence, JACM 10(3), 1963",
     "Origin of operator-precedence parsing; the idea behind %left / %right declarations."),
    ("[6]", "D. E. Knuth", "On the translation of languages from left to right, Information and Control 8(6), 1965",
     "Introduced LR(k) parsing, the theory under every Yacc-style generator."),
    ("[7]", "A. V. Aho, S. C. Johnson", "LR Parsing, ACM Computing Surveys 6(2), 1974",
     "Survey of LR/LALR construction and conflict handling."),
    ("[8]", "A. V. Aho, S. C. Johnson, J. D. Ullman", "Deterministic parsing of ambiguous grammars, CACM 18(8), 1975",
     "Resolving shift/reduce conflicts with precedence and associativity - how our ambiguous expr rule is made deterministic."),
    ("[9]", "K. D. Cooper, L. Torczon", "Engineering a Compiler, 2nd ed., Morgan Kaufmann, 2011",
     "IR design, constant folding and type conversion (planned phases)."),
    ("[10]", "A. W. Appel", "Modern Compiler Implementation in C, Cambridge Univ. Press, 1998",
     "Abstract syntax trees and symbol-table design in C."),
    ("[11]", "T. J. Parr, R. W. Quong", "ANTLR: A predicated-LL(k) parser generator, Software: Practice and Experience 25(7), 1995",
     "Alternative top-down generator; compared with LALR(1) Bison."),
    ("[12]", "C. Lattner, V. Adve", "LLVM: A compilation framework for lifelong program analysis & transformation, CGO 2004",
     "Modern three-address-style SSA IR; shows where TAC leads."),
]

# ---------------------------------------------------------------- DOCX
def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)

def code(doc, text, size=8):
    for line in text.split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.left_indent = Inches(0.2)
        r = p.add_run(line if line else " "); r.font.name = "Consolas"; r.font.size = Pt(size)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    doc.add_paragraph()

def table(doc, header, rows, widths=None, size=9):
    t = doc.add_table(rows=1, cols=len(header)); t.style = "Table Grid"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(size)
        r.font.color.rgb = RGBColor(255, 255, 255); shade(c, "1F3A5F")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""; r = cells[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(size)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Inches(w)
    doc.add_paragraph()

def para(doc, text, style=None):
    p = doc.add_paragraph(text, style=style); p.paragraph_format.space_after = Pt(6); return p

def build_docx():
    d = Document()
    st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    for s in d.sections:
        s.left_margin = s.right_margin = Inches(1); s.top_margin = s.bottom_margin = Inches(0.9)

    for _ in range(5): d.add_paragraph()
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(TITLE); r.bold = True; r.font.size = Pt(24); r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    for t_, sz in (("Project Report - Review 1", 14), (AUTHOR, 11), (COURSE, 11)):
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; rr = p.add_run(t_); rr.font.size = Pt(sz)
    d.add_page_break()

    d.add_heading("Abstract", 1)
    para(d, "This project builds a compiler for a small arithmetic expression language using Flex and Bison. The full "
            "pipeline is: tokenise the source, construct the parse tree, perform semantic validation, and generate Three "
            "Address Code (TAC). This Review-1 report covers the first milestone, about 25% of the work: the language "
            "definition, the Flex lexer that produces the token stream, and the Bison parser that resolves operator "
            "precedence and associativity, recovers from syntax errors, and constructs the parse tree. Semantic validation "
            "and TAC generation are the planned next stages.")

    d.add_heading("1. Problem Identification", 1)
    d.add_heading("1.1 Problem statement", 2)
    para(d, "Evaluating or translating an arithmetic expression such as  a + b * 2 - 4 / 2  looks trivial, but doing it "
            "correctly requires the classical compiler phases: recognising lexemes, respecting operator precedence and "
            "associativity, rejecting ill-formed or ill-typed programs, and producing a machine-independent intermediate form. "
            "Hand-written ad-hoc evaluators typically fail in four ways:")
    for b in ("Precedence and associativity are encoded in scattered if/else logic and are easy to get wrong "
              "(for example, 2^3^2 must be 2^9 = 512, not 8^2 = 64).",
              "Errors are reported one at a time, or not at all - an undeclared variable or a division by a constant zero "
              "surfaces only at run time.",
              "Mixed int/float arithmetic silently changes meaning unless conversions are made explicit.",
              "There is no intermediate representation, so no optimisation or retargeting is possible."):
        d.add_paragraph(b, style="List Bullet")
    d.add_heading("1.2 Objectives", 2)
    for b in ("O1  Specify a small, unambiguous arithmetic language: int/float declarations, assignment, print, operators + - * / % ^, unary minus, parentheses.   [done]",
              "O2  Lexical analysis: generate a token stream with Flex, with line numbers and error reporting for illegal characters.   [done]",
              "O3  Syntax analysis: write an LALR(1) grammar in Bison, resolve ambiguity through precedence declarations, and construct the parse tree.   [done]",
              "O4  Semantic validation: symbol table, type inference/checking and compile-time diagnostics.   [planned]",
              "O5  Intermediate code: generate Three Address Code with temporaries, explicit type conversions and constant folding.   [planned]"):
        d.add_paragraph(b)
    d.add_heading("1.3 Scope", 2)
    para(d, "In scope: one global scope, two types (int, float), the operators listed above. Out of scope: control flow, "
            "functions, arrays, strings, target-machine code generation and register allocation.")

    d.add_heading("2. Literature Survey", 1)
    para(d, "The design rests on four bodies of work: the phase structure of compilers, lexer/parser generator theory, "
            "resolution of expression-grammar ambiguity, and intermediate representations.")
    table(d, ["Ref", "Authors", "Work", "Relevance to this project"], LIT, widths=[0.4, 1.3, 2.3, 2.5], size=8)
    d.add_heading("2.1 Analysis and gap", 2)
    para(d, "Lex [2] and Yacc [3] established the generator approach that Flex and Bison still follow: regular expressions "
            "become a DFA scanner and a context-free grammar becomes an LALR(1) table. Floyd [5] and Aho-Johnson-Ullman [8] "
            "show why a deliberately ambiguous rule such as expr -> expr op expr is acceptable: precedence and "
            "associativity declarations turn shift/reduce conflicts into deterministic choices, giving a smaller grammar "
            "than the stratified expr/term/factor form. Top-down generators such as ANTLR [11] avoid left-recursion "
            "restrictions differently, while LLVM [12] shows the end point of the IR line that TAC begins.")
    para(d, "Textbook treatments [1][9][10] present the phases separately and usually stop at the parse tree or at "
            "TAC listings. This project aims at a small, integrated pipeline in which every phase can be run and observed on "
            "the same input; Review 1 delivers the front half of it.")
    para(d, "Note: citations should be re-checked against the library copies before final submission.")

    d.add_heading("3. Proposed System and Design", 1)
    d.add_heading("3.1 Architecture", 2)
    code(d, "source.ac --> [Flex lexer] --tokens--> [Bison parser] --parse tree--> [Semantic analyser] --> [TAC generator] --> TAC\n"
            "                 |                         |                                  |                       |\n"
            "          illegal-char errors       syntax errors +                    (planned)               (planned)\n"
            "                                    panic-mode recovery\n"
            "|<------------ Review 1: implemented ------------>|", size=7)
    d.add_heading("3.2 Language grammar", 2)
    code(d, GRAMMAR, 9)
    d.add_heading("3.3 Lexical specification (Flex)", 2)
    table(d, ["Token class", "Pattern / examples", "Purpose"], TOKENS, widths=[1.8, 2.6, 2.1])
    para(d, "Flex uses the longest-match rule, and keyword rules are listed before the identifier rule so that 'int' "
            "is a KEYWORD and not an IDENTIFIER. %option yylineno supplies line numbers. A character that matches no rule "
            "is reported as a lexical error and scanning continues.")
    d.add_heading("3.4 Parser and parse-tree construction (Bison)", 2)
    para(d, "Each grammar action calls a node constructor (mk_bin, mk_int, mk_decl ...) so that a syntax tree is "
            "assembled bottom-up during parsing. The %union carries integer, float, string and node values. The rule "
            "'error ;' implements panic-mode recovery: after a syntax error the parser discards tokens up to the next "
            "semicolon and continues, so one run reports all errors in the file.")
    para(d, "Ambiguity is resolved with precedence declarations (%left '+' '-', then %left '*' '/' '%', "
            "%precedence UMINUS, %right '^'). Making '^' bind tighter than unary minus follows mathematical convention (-2^2 = -4). "
            "Bison builds the grammar with no warnings and no conflicts.")
    d.add_heading("3.5 Planned phases (next reviews)", 2)
    table(d, ["Phase", "Planned design"], [
        ("Semantic validation", "Symbol table (name, type, initialised); type inference with int-to-float promotion; errors for undeclared variable, redeclaration, % on floats, constant division by zero; warnings for use before assignment and float-to-int narrowing."),
        ("Three Address Code", "At most one operator per instruction; temporaries _t1, _t2 ...; explicit inttofloat / floattoint; constant folding of literal-only operations.")],
        widths=[1.6, 4.9])

    d.add_heading("4. Implementation Status", 1)
    para(d, "Implementation is in C with Flex 2.6.4, Bison 3.8.2 and GCC. The table maps each Review-1 module to its source file.")
    table(d, ["Module", "File", "State"], [
        ("Lexical analyser (token stream)", "src/lexer.l, src/ast.c (token log)", "Done - tested"),
        ("Parser (grammar, precedence, recovery)", "src/parser.y", "Done - tested"),
        ("Parse-tree construction and printer", "src/ast.c", "Done - tested"),
        ("Driver", "src/main.c", "Done"),
        ("Regression suite", "run_tests.sh", f"{n_pass} / {n_pass + n_fail} passing"),
        ("Semantic analysis", "-", "Not started (planned)"),
        ("Three Address Code", "-", "Not started (planned)")], widths=[2.6, 2.4, 1.5])
    d.add_heading("4.1 Build and run", 2)
    code(d, "sh build.sh                       # bison -d, flex, gcc -> arithc.exe\n"
            "./arithc.exe prog.ac              # diagnostics, token stream, parse tree\n"
            "./arithc.exe -q prog.ac           # quiet: hide the token table\n"
            "sh run_tests.sh                   # regression suite", 9)

    d.add_heading("5. Results", 1)
    d.add_heading("5.1 Token stream", 2)
    d.add_paragraph("Source (tests/t4_demo.ac):"); code(d, demo_src, 9)
    d.add_paragraph("Output of the lexer (first 26 lines):"); code(d, "  LINE  TOKEN            LEXEME\n" + "\n".join(demo_tokens[1:26]), 8)
    d.add_heading("5.2 Parse tree", 2)
    d.add_paragraph("Parse tree for the same program. '*' binds tighter than '+', and '-' is left-associative, so "
                    "int c = a + b * 2 - 4 / 2 parses as (a + (b*2)) - (4/2):")
    code(d, demo_tree, 8)
    d.add_heading("5.3 Associativity of ^", 2)
    d.add_paragraph("For  int p = 2 ^ 3 ^ 2;  the second ^ is the right child of the first, i.e. 2^(3^2) = 512, as required:")
    code(d, pow_tree, 9)
    d.add_heading("5.4 Lexical and syntax error recovery", 2)
    d.add_paragraph("Source (tests/t3_syntax.ac):"); code(d, bad_src, 9)
    d.add_paragraph("Diagnostics: every bad line is reported and parsing continues to the end of the file."); code(d, bad_diag, 9)
    d.add_heading("5.5 Regression suite", 2)
    code(d, tests_out, 9)

    d.add_heading("6. Project Plan", 1)
    table(d, ["Phase", "Work", "Status"], PLAN, widths=[0.6, 4.4, 1.5])

    d.add_heading("7. Conclusion", 1)
    para(d, "Review 1 delivers the language definition, a Flex lexer producing a line-numbered token stream, and a Bison "
            "parser that builds the parse tree, resolves precedence and associativity through declarations, and recovers "
            "from syntax errors so that all errors are reported in one run. The remaining phases - semantic validation and "
            "Three Address Code generation - build directly on the parse tree produced here.")

    d.add_heading("References", 1)
    for a, b, c, _ in LIT:
        d.add_paragraph(f"{a} {b}, {c}.")
    d.add_heading("Appendix A - Source listing", 1)
    for f in ("src/lexer.l", "src/parser.y", "src/ast.h", "src/main.c"):
        d.add_heading(f, 2); code(d, (ROOT / f).read_text(encoding="utf-8"), 7)
    d.save(OUT / "REPORT.docx")

# ---------------------------------------------------------------- PPTX
NAVY, TEAL, GREY, LIGHT = PC(0x1F, 0x3A, 0x5F), PC(0x1B, 0x8A, 0x8A), PC(0x55, 0x5B, 0x66), PC(0xF2, 0xF5, 0xF9)
MUTE = PC(0xB5, 0xBD, 0xC9)

def build_pptx():
    prs = Presentation(); prs.slide_width = PI(13.333); prs.slide_height = PI(7.5)
    blank = prs.slide_layouts[6]

    def base(title, notes=""):
        s = prs.slides.add_slide(blank)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, PI(1.0)); bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
        tb = s.shapes.add_textbox(PI(0.5), PI(0.15), PI(12.3), PI(0.7)); tf = tb.text_frame
        p = tf.paragraphs[0]; p.text = title; p.font.size = PP(30); p.font.bold = True; p.font.color.rgb = PC(255, 255, 255)
        n = s.shapes.add_textbox(PI(12.3), PI(7.0), PI(0.9), PI(0.4)); n.text_frame.text = str(len(prs.slides)); n.text_frame.paragraphs[0].font.size = PP(12); n.text_frame.paragraphs[0].font.color.rgb = GREY
        if notes: s.notes_slide.notes_text_frame.text = notes
        return s

    def bullets(s, items, x=0.6, y=1.3, w=12.1, h=5.5, size=20):
        tb = s.shapes.add_textbox(PI(x), PI(y), PI(w), PI(h)); tf = tb.text_frame; tf.word_wrap = True
        for i, it in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            sub = it.startswith("  "); p.text = ("- " if sub else "") + it.strip(); p.level = 1 if sub else 0
            p.font.size = PP(size - 3 if sub else size); p.font.color.rgb = GREY if sub else PC(0x22, 0x22, 0x22); p.space_after = PP(8)

    def mono(s, text, x, y, w, h, size=12, fill=LIGHT):
        box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, PI(x), PI(y), PI(w), PI(h)); box.fill.solid(); box.fill.fore_color.rgb = fill; box.line.color.rgb = PC(0xCC, 0xD3, 0xDD)
        tf = box.text_frame; tf.word_wrap = False; tf.margin_left = PI(0.15); tf.margin_top = PI(0.1)
        for i, line in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.text = line if line else " "
            p.font.name = "Consolas"; p.font.size = PP(size); p.font.color.rgb = PC(0x1A, 0x1A, 0x1A); p.alignment = PP_ALIGN.LEFT

    def tbl(s, header, rows, x, y, w, colw, size=13, rowh=0.45):
        shape = s.shapes.add_table(len(rows) + 1, len(header), PI(x), PI(y), PI(w), PI(rowh * (len(rows) + 1)))
        t = shape.table
        for i, cw in enumerate(colw): t.columns[i].width = PI(cw)
        for i, h in enumerate(header):
            c = t.cell(0, i); c.text = h; c.fill.solid(); c.fill.fore_color.rgb = NAVY
            for p in c.text_frame.paragraphs: p.font.size = PP(size); p.font.bold = True; p.font.color.rgb = PC(255, 255, 255)
        for r, row in enumerate(rows, 1):
            for i, v in enumerate(row):
                c = t.cell(r, i); c.text = str(v); c.fill.solid(); c.fill.fore_color.rgb = LIGHT if r % 2 else PC(255, 255, 255)
                for p in c.text_frame.paragraphs: p.font.size = PP(size); p.font.color.rgb = PC(0x22, 0x22, 0x22)

    # 1 title
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height); bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
    tb = s.shapes.add_textbox(PI(0.8), PI(2.2), PI(11.7), PI(2)); tf = tb.text_frame; tf.word_wrap = True
    tf.text = "Compiler for an Arithmetic Expression Language"; tf.paragraphs[0].font.size = PP(42); tf.paragraphs[0].font.bold = True; tf.paragraphs[0].font.color.rgb = PC(255, 255, 255)
    p = tf.add_paragraph(); p.text = "using Lex/Flex and Yacc/Bison"; p.font.size = PP(28); p.font.color.rgb = PC(0x9F, 0xD8, 0xD8)
    p = tf.add_paragraph(); p.text = "Tokens  ->  Parse tree  ->  Semantic validation  ->  Three Address Code"; p.font.size = PP(20); p.font.color.rgb = PC(0xDD, 0xE6, 0xF0)
    tb = s.shapes.add_textbox(PI(0.8), PI(5.6), PI(11.7), PI(1.2)); tb.text_frame.word_wrap = True
    tb.text_frame.text = "Review 1   |   " + AUTHOR; tb.text_frame.paragraphs[0].font.size = PP(16); tb.text_frame.paragraphs[0].font.color.rgb = PC(255, 255, 255)
    p = tb.text_frame.add_paragraph(); p.text = COURSE; p.font.size = PP(16); p.font.color.rgb = PC(0xDD, 0xE6, 0xF0)

    # 2 problem
    s = base("Problem Identification", "Open with the 2^3^2 example: most people say 64, correct is 512. This is why precedence must come from a grammar, not ad-hoc code.")
    bullets(s, ["Evaluating  a + b * 2 - 4 / 2  correctly needs the compiler phases",
                "Ad-hoc evaluators break in four ways:",
                "  Precedence/associativity scattered in if-else:  2^3^2 must be 512, not 64",
                "  Errors found late or never: undeclared names, constant divide-by-zero",
                "  Silent int/float mixing changes meaning",
                "  No intermediate form -> no optimisation, no retargeting",
                "Need: a principled pipeline that rejects bad programs and emits machine-independent code"])

    # 3 objectives
    s = base("Objectives")
    tbl(s, ["#", "Objective", "Tool", "Status"], [
        ("O1", "Define a small arithmetic language (int, float, + - * / % ^, unary -, parentheses)", "Grammar", "Done"),
        ("O2", "Tokenise source with line numbers and lexical error reporting", "Flex", "Done"),
        ("O3", "LALR(1) grammar, precedence resolution, build the parse tree", "Bison", "Done"),
        ("O4", "Symbol table, type inference, compile-time diagnostics", "C", "Planned"),
        ("O5", "Three Address Code with temporaries, conversions, constant folding", "C", "Planned")],
        0.6, 1.4, 12.1, [0.8, 7.6, 1.6, 2.1], size=16, rowh=0.75)

    # 4-5 literature
    s = base("Literature Survey (1/2): Tools and Theory", "Lex/Yacc are the 1975 origin of Flex/Bison. Floyd and Aho-Johnson-Ullman justify the ambiguous expr rule plus precedence.")
    tbl(s, ["Ref", "Work", "Contribution used here"], [
        ("[1]", "Aho, Lam, Sethi, Ullman - Compilers (Dragon Book), 2006", "Phase structure, SDT, TAC, type rules"),
        ("[2]", "Lesk & Schmidt - Lex, 1975", "Regex -> DFA scanner (Flex)"),
        ("[3]", "Johnson - Yacc, 1975", "LALR(1) generator + actions (Bison)"),
        ("[4]", "Levine - flex & bison, 2009", "Lexer/parser integration, error recovery"),
        ("[5]", "Floyd - Operator precedence, 1963", "Basis of %left / %right"),
        ("[6]", "Knuth - LR(k) translation, 1965", "Theory of bottom-up parsing")],
        0.6, 1.4, 12.1, [0.8, 6.3, 5.0], size=15, rowh=0.7)

    s = base("Literature Survey (2/2): Resolution, IR and Gap")
    tbl(s, ["Ref", "Work", "Contribution used here"], [
        ("[7]", "Aho & Johnson - LR Parsing, 1974", "LALR construction, conflicts"),
        ("[8]", "Aho, Johnson, Ullman - Ambiguous grammars, 1975", "Precedence resolves shift/reduce conflicts"),
        ("[9]", "Cooper & Torczon - Engineering a Compiler, 2011", "IR design, folding, conversions"),
        ("[10]", "Appel - Modern Compiler Impl. in C, 1998", "AST + symbol table in C"),
        ("[11]", "Parr & Quong - ANTLR, 1995", "Top-down LL(k) alternative"),
        ("[12]", "Lattner & Adve - LLVM, 2004", "Where TAC/SSA IR leads")],
        0.6, 1.4, 12.1, [0.8, 6.3, 5.0], size=15, rowh=0.62)
    bullets(s, ["Gap: textbooks show phases separately. Goal here: one small pipeline where every phase can be run and observed."], y=6.0, h=1.0, size=18)

    # 6 language
    s = base("Language Specification")
    mono(s, GRAMMAR, 0.6, 1.3, 8.2, 5.6, size=13)
    bullets(s, ["Types: int, float", "Statements: declare, assign, print", "Comments: // and /* */", "Unary +/-, parentheses",
                "Ambiguous expr rule + precedence declarations = compact grammar"], x=9.0, y=1.3, w=4.0, h=5.5, size=16)

    # 7 architecture
    s = base("Compiler Architecture", "Walk left to right. The first three boxes are what is implemented for Review 1; the grey ones are the next stages.")
    stages = [("Source\n.ac", GREY), ("Flex\nLexer", TEAL), ("Bison\nParser", TEAL), ("Semantic\nAnalyser", MUTE), ("TAC\nGenerator", MUTE)]
    for i, (t_, col) in enumerate(stages):
        b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PI(0.6 + i * 2.5), PI(2.3), PI(2.0), PI(1.2)); b.fill.solid(); b.fill.fore_color.rgb = col; b.line.fill.background()
        b.text_frame.text = t_
        for p in b.text_frame.paragraphs: p.font.size = PP(17); p.font.bold = True; p.font.color.rgb = PC(255, 255, 255); p.alignment = PP_ALIGN.CENTER
        if i < len(stages) - 1:
            a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, PI(2.65 + i * 2.5), PI(2.7), PI(0.4), PI(0.35)); a.fill.solid(); a.fill.fore_color.rgb = NAVY; a.line.fill.background()
    labels = ["", "tokens", "parse tree", "typed tree", "TAC"]
    for i, l in enumerate(labels):
        if l:
            t = s.shapes.add_textbox(PI(0.5 + i * 2.5), PI(3.6), PI(2.2), PI(0.5)); t.text_frame.text = l
            for p in t.text_frame.paragraphs: p.font.size = PP(14); p.font.color.rgb = GREY; p.alignment = PP_ALIGN.CENTER
    t = s.shapes.add_textbox(PI(0.6), PI(4.3), PI(7.0), PI(0.5)); t.text_frame.text = "Teal = implemented for Review 1      Grey = planned"
    t.text_frame.paragraphs[0].font.size = PP(16); t.text_frame.paragraphs[0].font.color.rgb = GREY
    bullets(s, ["Errors per phase: illegal character | syntax error (panic-mode recovery) | undeclared / type errors (planned)"], y=5.2, size=18)

    # 8 lexer
    s = base("Phase 1: Lexical Analysis (Flex)")
    tbl(s, ["Token class", "Pattern / examples", "Purpose"], TOKENS, 0.6, 1.3, 12.1, [3.6, 5.0, 3.5], size=13, rowh=0.5)
    bullets(s, ["Longest-match rule; keywords listed before IDENTIFIER; %option yylineno for line numbers",
                "Unknown character -> lexical error, scanning continues"], y=5.9, h=1.2, size=16)

    # 9 token demo
    s = base("Demo: Token Stream", "Live: ./arithc.exe tests/t4_demo.ac. Each chip is one token; colour = token class; the number on the left is the source line.")
    toks = []
    for l in demo_tokens[1:-1]:
        parts = l.split(None, 2)
        if len(parts) == 3 and parts[0].isdigit(): toks.append((parts[1], parts[2], int(parts[0])))
    draw_token_chips(s, toks, 0.6, 1.4, 12.1)
    t = s.shapes.add_textbox(PI(0.6), PI(6.5), PI(12), PI(0.5)); t.text_frame.text = f"{len(toks)} tokens from 7 source lines (tests/t4_demo.ac)"
    t.text_frame.paragraphs[0].font.size = PP(15); t.text_frame.paragraphs[0].font.color.rgb = GREY

    # 10 parser
    s = base("Phase 2: Syntax Analysis (Bison)")
    bullets(s, ["LALR(1) grammar; actions call mk_bin / mk_int / mk_decl -> tree built bottom-up",
                "Precedence table resolves every shift/reduce conflict:",
                "  %left + -   <   %left * / %   <   %precedence UMINUS   <   %right ^",
                "  Right-assoc ^:  2^3^2 = 2^(3^2) = 512;    -2^2 = -(2^2) = -4",
                "Panic-mode recovery:  stmt -> error ';'   reports all errors in one run",
                "Bison builds with zero warnings and zero conflicts"], size=19)

    # 11 parse tree (drawn)
    s = base("Demo: Parse Tree", "Read bottom-up: b*2 is evaluated first, then a+(b*2), 4/2 on the other side, and the subtraction last.")
    roots = parse_statements(demo_tree.split("\n")[1:])
    c_root = next(r for r in roots if r["label"] == "Decl int c")
    t = s.shapes.add_textbox(PI(0.6), PI(1.15), PI(8), PI(0.5)); t.text_frame.text = "int c = a + b * 2 - 4 / 2;"
    t.text_frame.paragraphs[0].font.name = "Consolas"; t.text_frame.paragraphs[0].font.size = PP(20); t.text_frame.paragraphs[0].font.bold = True
    draw_tree(s, c_root, 0.6, 1.9, 8.0, 4.9, font=14)
    bullets(s, ["* binds tighter than +, so b*2 is a subtree of +",
                "- is left-associative: (a + b*2) - (4/2)",
                "Leaves are tokens: variables and constants",
                "Built by Bison actions while parsing"], x=9.0, y=1.5, w=4.1, h=5.3, size=16)

    # 12 ^ assoc + errors
    s = base("Demo: Associativity and Error Recovery", "Left: 2^3^2 nests to the right. Right: three bad lines, all reported, parsing continues.")
    proots = parse_statements(prec_tree.split("\n")[1:])
    p_root = next(r for r in proots if r["label"] == "Decl int p")
    t = s.shapes.add_textbox(PI(0.6), PI(1.15), PI(5.5), PI(0.5)); t.text_frame.text = "int p = 2 ^ 3 ^ 2;"
    t.text_frame.paragraphs[0].font.name = "Consolas"; t.text_frame.paragraphs[0].font.size = PP(20); t.text_frame.paragraphs[0].font.bold = True
    draw_tree(s, p_root, 0.6, 1.9, 5.4, 3.0, font=14)
    bullets(s, ["Second ^ is the RIGHT child: 2^(3^2) = 512"], x=0.6, y=5.2, w=5.6, h=1.2, size=17)
    mono(s, bad_src, 6.5, 1.3, 6.2, 2.4, size=13)
    mono(s, bad_diag, 6.5, 3.9, 6.2, 2.2, size=11)

    # 13 status
    s = base("Review 1 Status and Plan", "25% milestone = language + lexer + grammar + parse tree. Everything after phase 3 is planned, not built.")
    tbl(s, ["Phase", "Work", "Status"], PLAN, 0.6, 1.4, 12.1, [0.9, 8.3, 2.9], size=15, rowh=0.7)
    bullets(s, [f"Regression suite: {n_pass}/{n_pass + n_fail} tests pass"], y=6.4, h=0.6, size=16)

    # 14 next
    s = base("Next Stage: Semantic Validation and TAC (planned)")
    bullets(s, ["Semantic validation",
                "  Symbol table; type inference with int -> float promotion",
                "  Errors: undeclared variable, redeclaration, % on floats, constant division by zero",
                "  Warnings: use before assignment, float -> int narrowing",
                "Three Address Code",
                "  <= 1 operator per instruction, temporaries _t1, _t2 ...",
                "  Explicit inttofloat / floattoint; constant folding"], size=20)

    # 15 conclusion
    s = base("Conclusion")
    bullets(s, ["Review 1 delivered: language definition, Flex lexer, Bison parser, parse tree, error recovery",
                "Precedence and right-associative ^ come from the grammar, not ad-hoc code",
                f"{n_pass}/{n_pass + n_fail} regression tests pass",
                "Next: semantic validation, then Three Address Code generation"], size=20)

    # 16 Q&A
    s = base("Questions?", "Likely questions: why LALR over LL; why is ^ above unary minus; how does Bison resolve conflicts; why TAC; what does panic mode do.")
    bullets(s, ["Why Bison (LALR(1)) instead of a hand-written recursive-descent parser?",
                "How does an ambiguous expr rule stay deterministic?   -> precedence/associativity",
                "Why Three Address Code?   -> machine independent, optimisable, one operator per line",
                "What does 'error ;' do?   -> panic-mode: skip to the next ';' and resume"], size=20)
    prs.save(OUT / "PRESENTATION.pptx")

if __name__ == "__main__":
    build_docx(); build_pptx(); print("ok")
